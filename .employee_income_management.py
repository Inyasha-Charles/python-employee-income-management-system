from pathlib import Path
from openpyxl import Workbook, load_workbook


# ==========================================================
# EXCEL FILE
# ==========================================================

FILE_NAME = "Users.xlsx"


# ==========================================================
# COLUMN HEADINGS
# ==========================================================

HEADERS = [
    "Username",
    "Password",
    "Role",
    "Payment Type",
    "Salary",
    "Wage",
    "Tips",
    "Total Income"
]


# ==========================================================
# OPEN OR CREATE EXCEL FILE
# ==========================================================

def open_users():

    if Path(FILE_NAME).exists():

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

    else:

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Users"

        sheet.append(HEADERS)

        workbook.save(FILE_NAME)

    return workbook, sheet


# ==========================================================
# FIND USER'S ROW
# ==========================================================

def find_user(sheet, username):

    for row in range(2, sheet.max_row + 1):

        saved_username = sheet.cell(row, 1).value

        if (
            saved_username is not None
            and str(saved_username).strip().lower()
            == username.strip().lower()
        ):

            return row

    return None


# ==========================================================
# FIRST ADMIN SETUP
# ==========================================================

def setup_first_admin():

    workbook, sheet = open_users()

    admin_exists = False

    for row in range(2, sheet.max_row + 1):

        role = sheet.cell(row, 3).value

        if (
            role is not None
            and str(role).strip().lower() == "admin"
        ):

            admin_exists = True
            break

    workbook.close()

    if admin_exists:
        return

    print("\n================================")
    print("       FIRST ADMIN SETUP")
    print("================================")
    print("No administrator account exists.")
    print("You must create the first admin.")
    print("================================\n")

    while True:

        username = input("Create admin username: ").strip()

        if not username:

            print("Username cannot be empty.\n")
            continue

        password = input("Create admin password: ")

        if not password:

            print("Password cannot be empty.\n")
            continue

        confirm_password = input("Confirm admin password: ")

        if password != confirm_password:

            print("Passwords do not match. Try again.\n")
            continue

        workbook, sheet = open_users()

        if find_user(sheet, username) is not None:

            print("That username already exists.\n")
            workbook.close()
            continue

        # Create the first admin
        sheet.append([
            username,
            password,
            "admin",
            "s",
            0,
            0,
            0,
            0
        ])

        workbook.save(FILE_NAME)
        workbook.close()

        print("\nFirst admin account created successfully!")
        print("Username:", username)
        print("Role: admin\n")

        break


# ==========================================================
# REGISTER NEW USER
# ADMIN ONLY
# ==========================================================

def register():

    print("\n================================")
    print("        REGISTER USER")
    print("================================\n")

    username = input("Create username: ").strip()
    password = input("Create password: ")

    if not username or not password:

        print("Username and password cannot be empty.\n")
        return

    workbook, sheet = open_users()

    if find_user(sheet, username) is not None:

        print("That username already exists.\n")
        workbook.close()
        return

    # ------------------------------------------------------
    # ROLE
    # ------------------------------------------------------

    print("\nChoose user role:")
    print("1. Admin")
    print("2. HR")
    print("3. Normal user")

    role_choice = input("Choose a role: ")

    if role_choice == "1":

        role = "admin"

    elif role_choice == "2":

        role = "hr"

    elif role_choice == "3":

        role = "normal"

    else:

        print("Invalid role.")
        workbook.close()
        return

    # ------------------------------------------------------
    # PAYMENT TYPE
    # ------------------------------------------------------

    payment_type = input(
        "\nWage or salary? w/s: "
    ).lower()

    if payment_type not in ("w", "s"):

        print("Please enter w or s.\n")
        workbook.close()
        return

    try:

        # --------------------------------------------------
        # SALARY / WAGE
        # --------------------------------------------------

        if payment_type == "s":

            salary = float(input("Enter salary: "))
            wage = 0

        else:

            wage = float(input("Enter wage: "))
            salary = 0

        if salary < 0 or wage < 0:

            print("Income cannot be negative.\n")
            workbook.close()
            return

        # --------------------------------------------------
        # TIPS
        # --------------------------------------------------

        answer = input(
            "Do you receive tips? y/n: "
        ).lower()

        if answer == "y":

            tips = float(input("Enter tips: "))

        elif answer == "n":

            tips = 0

        else:

            print("Please enter y or n.\n")
            workbook.close()
            return

        if tips < 0:

            print("Tips cannot be negative.\n")
            workbook.close()
            return

        # --------------------------------------------------
        # TOTAL INCOME
        # --------------------------------------------------

        if payment_type == "s":

            total_income = salary + tips

        else:

            total_income = wage + tips

    except ValueError:

        print("Please enter valid numbers.\n")
        workbook.close()
        return

    # ------------------------------------------------------
    # ADD USER
    # ------------------------------------------------------

    sheet.append([
        username,
        password,
        role,
        payment_type,
        salary,
        wage,
        tips,
        total_income
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nRegistration successful!")
    print("Username:", username)
    print("Role:", role)
    print()


# ==========================================================
# LOGIN
# ==========================================================

def login():

    print("\n================================")
    print("             LOGIN")
    print("================================\n")

    username = input("Username: ").strip()
    password = input("Password: ")

    if not Path(FILE_NAME).exists():

        print("No registered users found.\n")
        return

    workbook, sheet = open_users()

    row = find_user(sheet, username)

    # Check username AND password
    if row is None or sheet.cell(row, 2).value != password:

        print("Invalid username or password.\n")
        workbook.close()
        return

    # Get role
    role = sheet.cell(row, 3).value

    workbook.close()

    print("\nLogin successful!")
    print("Welcome,", username)
    print("Role:", role)
    print()

    income_menu(username, role)


# ==========================================================
# MAIN USER MENU
# ==========================================================

def income_menu(username, role):

    while True:

        print("\n================================")
        print("          INCOME MENU")
        print("================================")

        print("1. Show my salary/wage")
        print("2. Show my tips")
        print("3. Add/Edit my tips")

        # HR + Admin
        if role in ("admin", "hr"):

            print("4. Edit profile")
            print("5. View all users")
            print("6. View specific user")

        # Admin only
        if role == "admin":

            print("7. Register new user")
            print("8. Delete user")

        print("9. Logout")
        print("================================")

        choice = input("Choose an option: ")
        print()

        # ==================================================
        # 1. SHOW OWN SALARY / WAGE
        # ==================================================

        if choice == "1":

            workbook, sheet = open_users()

            row = find_user(sheet, username)

            if row is None:

                workbook.close()
                print("Your account could not be found.")
                return

            payment_type = sheet.cell(row, 4).value
            salary = float(sheet.cell(row, 5).value or 0)
            wage = float(sheet.cell(row, 6).value or 0)

            workbook.close()

            if payment_type == "s":

                print("Your salary:", salary)

            else:

                print("Your wage:", wage)

        # ==================================================
        # 2. SHOW OWN TIPS
        # ==================================================

        elif choice == "2":

            workbook, sheet = open_users()

            row = find_user(sheet, username)

            if row is None:

                workbook.close()
                print("Your account could not be found.")
                return

            tips = float(sheet.cell(row, 7).value or 0)

            workbook.close()

            print("Your tips:", tips)

        # ==================================================
        # 3. EDIT OWN TIPS
        # ALL USERS
        # ==================================================

        elif choice == "3":

            edit_tips(username)

        # ==================================================
        # 4. EDIT PROFILE
        # ADMIN + HR
        # ==================================================

        elif choice == "4":

            if role in ("admin", "hr"):

                target_username = input(
                    "Enter username to edit: "
                ).strip()

                edit_user_profile(
                    target_username,
                    role
                )

            else:

                print(
                    "You do not have permission "
                    "to edit profiles."
                )

        # ==================================================
        # 5. VIEW ALL USERS
        # ADMIN + HR
        # ==================================================

        elif choice == "5":

            if role in ("admin", "hr"):

                view_all_users()

            else:

                print(
                    "You do not have permission "
                    "to view users."
                )

        # ==================================================
        # 6. VIEW SPECIFIC USER
        # ADMIN + HR
        # ==================================================

        elif choice == "6":

            if role in ("admin", "hr"):

                target_username = input(
                    "Enter username to view: "
                ).strip()

                view_user(target_username)

            else:

                print(
                    "You do not have permission "
                    "to view users."
                )

        # ==================================================
        # 7. REGISTER USER
        # ADMIN ONLY
        # ==================================================

        elif choice == "7":

            if role == "admin":

                register()

            else:

                print(
                    "You do not have permission "
                    "to register users."
                )

        # ==================================================
        # 8. DELETE USER
        # ADMIN ONLY
        # ==================================================

        elif choice == "8":

            if role == "admin":

                delete_user()

            else:

                print(
                    "You do not have permission "
                    "to delete users."
                )

        # ==================================================
        # 9. LOGOUT
        # ==================================================

        elif choice == "9":

            print("You have been logged out.\n")
            return

        else:

            print("Invalid option.")


# ==========================================================
# EDIT TIPS
# ALL USERS
# ==========================================================

def edit_tips(username):

    workbook, sheet = open_users()

    row = find_user(sheet, username)

    if row is None:

        print("User not found.\n")
        workbook.close()
        return

    current_tips = float(
        sheet.cell(row, 7).value or 0
    )

    print("Current tips:", current_tips)

    try:

        new_tips = float(
            input("Enter new tips amount: ")
        )

        if new_tips < 0:

            print("Tips cannot be negative.\n")
            workbook.close()
            return

        # Column 7 = Tips
        sheet.cell(row, 7).value = new_tips

        # Get salary/wage
        payment_type = sheet.cell(row, 4).value
        salary = float(sheet.cell(row, 5).value or 0)
        wage = float(sheet.cell(row, 6).value or 0)

        # Recalculate total
        if payment_type == "s":

            total_income = salary + new_tips

        else:

            total_income = wage + new_tips

        # Column 8 = Total Income
        sheet.cell(row, 8).value = total_income

        workbook.save(FILE_NAME)
        workbook.close()

        print("Tips updated successfully.")

    except ValueError:

        print("Please enter a valid number.")
        workbook.close()


# ==========================================================
# EDIT USER PROFILE
#
# ADMIN:
#   Can edit username, password, role,
#   payment type, salary, wage and tips.
#
# HR:
#   Can edit payment type, salary, wage and tips.
#   CANNOT edit username, password or role.
# ==========================================================

def edit_user_profile(username, editor_role):

    workbook, sheet = open_users()

    row = find_user(sheet, username)

    if row is None:

        print("User not found.\n")
        workbook.close()
        return

    print("\n================================")
    print("        EDIT USER PROFILE")
    print("================================")

    print("Editing user:", username)
    print()

    # ======================================================
    # ADMIN ONLY INFORMATION
    # ======================================================

    if editor_role == "admin":

        # --------------------------------------------------
        # USERNAME
        # --------------------------------------------------

        new_username = input(
            "New username "
            "(press Enter to keep current): "
        ).strip()

        if new_username == "":

            new_username = sheet.cell(row, 1).value

        else:

            existing_row = find_user(
                sheet,
                new_username
            )

            if (
                existing_row is not None
                and existing_row != row
            ):

                print("That username already exists.")
                workbook.close()
                return

        # --------------------------------------------------
        # PASSWORD
        # --------------------------------------------------

        new_password = input(
            "New password "
            "(press Enter to keep current): "
        )

        if new_password == "":

            new_password = sheet.cell(row, 2).value

        # --------------------------------------------------
        # ROLE
        # --------------------------------------------------

        print("\nAvailable roles:")
        print("1. Admin")
        print("2. HR")
        print("3. Normal")

        current_role = sheet.cell(row, 3).value

        print("Current role:", current_role)

        role_choice = input(
            "Choose new role "
            "(press Enter to keep current): "
        )

        if role_choice == "":

            new_role = current_role

        elif role_choice == "1":

            new_role = "admin"

        elif role_choice == "2":

            new_role = "hr"

        elif role_choice == "3":

            new_role = "normal"

        else:

            print("Invalid role.")
            workbook.close()
            return

    else:

        # --------------------------------------------------
        # HR CANNOT CHANGE THESE
        # --------------------------------------------------

        new_username = sheet.cell(row, 1).value
        new_password = sheet.cell(row, 2).value
        new_role = sheet.cell(row, 3).value

    # ======================================================
    # PAYMENT TYPE
    # ADMIN + HR
    # ======================================================

    current_payment = sheet.cell(row, 4).value

    payment_type = input(
        "\nWage or salary? w/s "
        f"(current: {current_payment}): "
    ).lower()

    if payment_type == "":

        payment_type = current_payment

    if payment_type not in ("w", "s"):

        print("Invalid payment type.")
        workbook.close()
        return

    # ======================================================
    # SALARY / WAGE
    # ADMIN + HR
    # ======================================================

    try:

        if payment_type == "s":

            current_salary = (
                sheet.cell(row, 5).value or 0
            )

            salary_input = input(
                f"New salary "
                f"(current: {current_salary}): "
            )

            if salary_input == "":

                salary = float(current_salary)

            else:

                salary = float(salary_input)

            wage = 0

        else:

            current_wage = (
                sheet.cell(row, 6).value or 0
            )

            wage_input = input(
                f"New wage "
                f"(current: {current_wage}): "
            )

            if wage_input == "":

                wage = float(current_wage)

            else:

                wage = float(wage_input)

            salary = 0

        # ==================================================
        # TIPS
        # ADMIN + HR
        # ==================================================

        current_tips = (
            sheet.cell(row, 7).value or 0
        )

        tips_input = input(
            f"New tips "
            f"(current: {current_tips}): "
        )

        if tips_input == "":

            tips = float(current_tips)

        else:

            tips = float(tips_input)

        # ==================================================
        # VALIDATE
        # ==================================================

        if salary < 0 or wage < 0 or tips < 0:

            print("Amounts cannot be negative.")
            workbook.close()
            return

        # ==================================================
        # TOTAL INCOME
        # ==================================================

        if payment_type == "s":

            total_income = salary + tips

        else:

            total_income = wage + tips

    except ValueError:

        print("Please enter valid numbers.")
        workbook.close()
        return

    # ======================================================
    # SAVE CHANGES
    # ======================================================

    sheet.cell(row, 1).value = new_username
    sheet.cell(row, 2).value = new_password
    sheet.cell(row, 3).value = new_role
    sheet.cell(row, 4).value = payment_type
    sheet.cell(row, 5).value = salary
    sheet.cell(row, 6).value = wage
    sheet.cell(row, 7).value = tips
    sheet.cell(row, 8).value = total_income

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nUser profile updated successfully.")


# ==========================================================
# VIEW ALL USERS
# ADMIN + HR
# ==========================================================

def view_all_users():

    print("\n================================")
    print("           ALL USERS")
    print("================================")

    workbook, sheet = open_users()

    if sheet.max_row < 2:

        print("There are no users registered.")
        workbook.close()
        return

    for row in range(2, sheet.max_row + 1):

        username = sheet.cell(row, 1).value
        role = sheet.cell(row, 3).value
        payment_type = sheet.cell(row, 4).value
        salary = sheet.cell(row, 5).value or 0
        wage = sheet.cell(row, 6).value or 0
        tips = sheet.cell(row, 7).value or 0
        total_income = sheet.cell(row, 8).value or 0

        print("--------------------------------")
        print("Username:", username)
        print("Role:", role)

        if payment_type == "s":

            print("Payment type: Salary")
            print("Salary:", salary)

        else:

            print("Payment type: Wage")
            print("Wage:", wage)

        print("Tips:", tips)
        print("Total income:", total_income)

    print("--------------------------------")

    workbook.close()


# ==========================================================
# VIEW SPECIFIC USER
# ADMIN + HR
# ==========================================================

def view_user(username):

    workbook, sheet = open_users()

    row = find_user(sheet, username)

    if row is None:

        print("User not found.")
        workbook.close()
        return

    username = sheet.cell(row, 1).value
    role = sheet.cell(row, 3).value
    payment_type = sheet.cell(row, 4).value
    salary = sheet.cell(row, 5).value or 0
    wage = sheet.cell(row, 6).value or 0
    tips = sheet.cell(row, 7).value or 0
    total_income = sheet.cell(row, 8).value or 0

    print("\n================================")
    print("          USER INFORMATION")
    print("================================")

    print("Username:", username)
    print("Role:", role)

    if payment_type == "s":

        print("Payment type: Salary")
        print("Salary:", salary)

    else:

        print("Payment type: Wage")
        print("Wage:", wage)

    print("Tips:", tips)
    print("Total income:", total_income)

    print("================================")

    workbook.close()


# ==========================================================
# DELETE USER
# ADMIN ONLY
# ==========================================================

def delete_user():

    print("\n================================")
    print("          DELETE USER")
    print("================================")

    username = input(
        "Enter the username you want to delete: "
    ).strip()

    workbook, sheet = open_users()

    row = find_user(sheet, username)

    if row is None:

        print("User not found.\n")
        workbook.close()
        return

    target_role = sheet.cell(row, 3).value

    print("\nUser found.")
    print("Username:", username)
    print("Role:", target_role)

    # Prevent deleting the only admin
    if target_role == "admin":

        admin_count = 0

        for current_row in range(
            2,
            sheet.max_row + 1
        ):

            current_role = sheet.cell(
                current_row, 3
            ).value

            if (
                current_role is not None
                and str(current_role).strip().lower()
                == "admin"
            ):

                admin_count += 1

        if admin_count <= 1:

            print(
                "\nYou cannot delete the only "
                "administrator account."
            )

            workbook.close()
            return

    confirmation = input(
        "\nAre you sure you want to "
        "delete this user? yes/no: "
    ).lower()

    if confirmation != "yes":

        print("Deletion cancelled.\n")
        workbook.close()
        return

    sheet.delete_rows(row, 1)

    workbook.save(FILE_NAME)
    workbook.close()

    print("User deleted successfully.\n")


# ==========================================================
# FIRST ADMIN SETUP
# ==========================================================

setup_first_admin()


# ==========================================================
# MAIN MENU
# ==========================================================

while True:

    print("\n================================")
    print("        INCOME SYSTEM")
    print("================================")
    print("1. Login")
    print("2. Exit")
    print("================================")

    choice = input("Choose an option: ")
    print()

    if choice == "1":

        login()

    elif choice == "2":

        print("Goodbye!")
        break

    else:

        print("Invalid option. Choose 1 or 2.\n")